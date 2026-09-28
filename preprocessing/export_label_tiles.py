#!/usr/bin/env python3
"""Export a blinded sample of tiles for pathologists to label.

The tissue compartments come from zero-shot CONCH labels that were never checked
against a pathologist tile by tile. This draws a stratified sample of tiles,
--per-class per predicted class and at most --per-slide-class tiles from any one
slide and class, from the slides in the modelled cohort, and writes:

- tiles/<id>.png: the 256 x 256 tile at the resolution it was classified, shown
  enlarged to --display px so it is easy to read. File names are random ids and
  carry no hint of the predicted class; the sample is shuffled.
- labelling_sheet.xlsx: one row per tile, the image embedded, and a drop-down
  of the cohort's classes plus "Other / not tissue" and "Unsure". Give one copy
  to each pathologist.
- key.csv: id -> slide, x, y, predicted class and probability. Keep this away
  from the pathologists until they have finished.

Afterwards, analysis/label_agreement.py turns the returned sheets into a
confusion matrix, per-class precision and recall, and agreement between the two
pathologists.

Reads the classifier's JSONL ({slide}_tiles.jsonl) and the WSIs with OpenSlide,
so it runs where both live.

Usage:

    python preprocessing/export_label_tiles.py --cohort blca \\
        --jsonl-dir /home/sorkwos/links/scratch/blca_jsons \\
        --svs-root /home/sorkwos/links/scratch/TCGA-BLCA-p2 \\
        --svs-root /home/sorkwos/links/scratch/TCGA-BLCA/WSI/output_folder1 \\
        --patients-file results/modelled_patients_blca.txt \\
        --out /scratch/$USER/label_tiles_blca
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import random
import uuid

OTHER = "Other / not tissue"
UNSURE = "Unsure"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cohort", required=True, choices=["blca", "brca"])
    ap.add_argument("--jsonl-dir", required=True)
    ap.add_argument("--svs-root", action="append", required=True)
    ap.add_argument("--patients-file", default=None, help="restrict to modelled patients")
    ap.add_argument("--per-class", type=int, default=25)
    ap.add_argument("--per-slide-class", type=int, default=2)
    ap.add_argument("--display", type=int, default=512)
    ap.add_argument("--seed", type=int, default=20260928)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    rng = random.Random(args.seed)

    keep = None
    if args.patients_file:
        keep = {l.strip()[:12] for l in open(args.patients_file) if l.strip()}
    svs = {}
    for root in args.svs_root:
        for p in glob.glob(os.path.join(root, "**", "*.svs"), recursive=True):
            svs[os.path.splitext(os.path.basename(p))[0]] = p
    paths = sorted(glob.glob(os.path.join(args.jsonl_dir, "*DX1*_tiles.jsonl")))
    rng.shuffle(paths)

    # reservoir per class, respecting the per-slide cap
    classes, pool = None, {}
    for jp in paths:
        slide = os.path.basename(jp).replace("_tiles.jsonl", "")
        if keep is not None and slide[:12] not in keep:
            continue
        if slide not in svs:
            continue
        by_class = {}
        with open(jp) as f:
            for ln in f:
                if not ln.strip():
                    continue
                r = json.loads(ln)
                if classes is None:
                    classes = list(r["all_scores"].keys())
                by_class.setdefault(r["pred_label"], []).append(r)
        for c, recs in by_class.items():
            pool.setdefault(c, [])
            for r in rng.sample(recs, min(args.per_slide_class, len(recs))):
                pool[c].append((slide, r))
        if classes and all(len(pool.get(c, [])) >= 4 * args.per_class for c in classes):
            break
    if classes is None:
        raise SystemExit("no JSONL tiles matched the slides found under --svs-root")

    sample = []
    for c in classes:
        cand = pool.get(c, [])
        sample += rng.sample(cand, min(args.per_class, len(cand)))
    rng.shuffle(sample)

    from openslide import OpenSlide
    os.makedirs(os.path.join(args.out, "tiles"), exist_ok=True)
    key_rows, open_slides = [], {}
    for slide, r in sample:
        tid = uuid.UUID(int=rng.getrandbits(128)).hex[:10]
        x, y, size = r["coords"]["x"], r["coords"]["y"], r["coords"].get("size", 256)
        sl = open_slides.get(slide) or open_slides.setdefault(slide, OpenSlide(svs[slide]))
        img = sl.read_region((x, y), 0, (size, size)).convert("RGB").resize((args.display, args.display))
        img.save(os.path.join(args.out, "tiles", f"{tid}.png"))
        key_rows.append({"id": tid, "slide": slide, "x": x, "y": y,
                         "predicted": r["pred_label"], "predicted_prob": r["pred_prob"]})

    with open(os.path.join(args.out, "key.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(key_rows[0].keys()))
        w.writeheader(); w.writerows(key_rows)

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(["id", "tile", "label", "comment"])
    opts = wb.create_sheet("options")
    for i, c in enumerate(classes + [OTHER, UNSURE], start=1):
        opts.cell(row=i, column=1, value=c)
    dv = DataValidation(type="list", formula1=f"=options!$A$1:$A${len(classes) + 2}", allow_blank=True)
    ws.add_data_validation(dv)
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 44
    for i, row in enumerate(key_rows, start=2):
        ws.cell(row=i, column=1, value=row["id"])
        im = XLImage(os.path.join(args.out, "tiles", f"{row['id']}.png"))
        im.width = im.height = 240
        ws.add_image(im, f"B{i}")
        ws.row_dimensions[i].height = 185
        dv.add(f"C{i}")
    wb.save(os.path.join(args.out, "labelling_sheet.xlsx"))
    counts = {c: sum(1 for k in key_rows if k["predicted"] == c) for c in classes}
    with open(os.path.join(args.out, "export_meta.json"), "w") as f:
        json.dump({"cohort": args.cohort, "n_tiles": len(key_rows), "per_predicted_class": counts,
                   "classes": classes, "extra_options": [OTHER, UNSURE], "seed": args.seed}, f, indent=2)
    print(f"wrote {len(key_rows)} tiles to {args.out}: {counts}")


if __name__ == "__main__":
    main()
